"""Leitura de planilha sob template — o motor S2 Data Intake.

Planilha é a única fonte do AlupData editada por uma pessoa. Coluna renomeada,
título acima do cabeçalho e número com vírgula são o normal, não a exceção.
Por isso o arquivo é conferido contra o template **antes** de qualquer linha ser
lida: arquivo fora do template é rejeitado inteiro, com erro dizendo o que falta
e o que sobra.

Desenho em `docs/arquitetura/s2-data-intake.md`.
"""

from __future__ import annotations

import csv
from dataclasses import dataclass, field
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import TYPE_CHECKING, Any

if TYPE_CHECKING:
    from collections.abc import Iterator


class TemplateInvalidoError(ValueError):
    """Arquivo não corresponde ao template declarado."""


@dataclass(frozen=True)
class TemplatePlanilha:
    """Contrato de uma planilha aceita pelo motor.

    `colunas` mapeia o cabeçalho como está no arquivo para o nome do campo no
    schema Pydantic da fonte — é ele que decide tipo e regra de negócio.
    """

    colunas: dict[str, str]
    delimitador: str = ";"  # Excel brasileiro exporta com ponto e vírgula
    encoding: str = "utf-8-sig"  # lê com e sem BOM
    linha_cabecalho: int = 1  # 1-based; > 1 quando há título acima
    aba: str | None = None  # None = primeira aba
    decimal_brasileiro: bool = True
    colunas_decimais: frozenset[str] = field(default_factory=frozenset)


def decimal_br(valor: str | None) -> Decimal | None:
    """Converte número em formato brasileiro (1.400,00) para Decimal.

    Campo vazio vira None — ausência não é zero. `,00` é zero de verdade.
    """
    if valor is None:
        return None
    texto = str(valor).strip().replace(".", "").replace(",", ".")
    if texto in {"", "."}:
        return None
    if texto.startswith("."):
        texto = "0" + texto
    try:
        return Decimal(texto)
    except InvalidOperation as exc:
        raise ValueError(f"número inválido: {valor!r}") from exc


def validar_cabecalho(encontrado: list[str], template: TemplatePlanilha) -> None:
    """Levanta `TemplateInvalidoError` quando o cabeçalho não bate com o template."""
    achadas = {c.strip() for c in encontrado if c and c.strip()}
    esperadas = set(template.colunas)
    faltam = sorted(esperadas - achadas)
    sobram = sorted(achadas - esperadas)
    if faltam or sobram:
        raise TemplateInvalidoError(f"planilha fora do template: faltam {faltam}; sobram {sobram}")


def ler_tabela(caminho: Path | str, template: TemplatePlanilha) -> Iterator[dict[str, Any]]:
    """Linhas da planilha já renomeadas para os campos do schema.

    Aceita `.csv` e `.xlsx`. Linha totalmente vazia é ignorada; célula vazia
    vira None, nunca string vazia.
    """
    caminho = Path(caminho)
    linhas = _linhas_xlsx(caminho, template) if caminho.suffix.lower() == ".xlsx" else _linhas_csv(caminho, template)

    cabecalho = next(linhas, None)
    if cabecalho is None:
        raise TemplateInvalidoError("planilha vazia: nenhum cabeçalho encontrado")
    validar_cabecalho(cabecalho, template)

    posicoes = {i: template.colunas[nome.strip()] for i, nome in enumerate(cabecalho) if nome and nome.strip()}
    for linha in linhas:
        if not any(_texto(c) for c in linha):
            continue  # linha em branco: resíduo de edição, não dado
        yield _montar(linha, posicoes, template)


def _montar(linha: list[Any], posicoes: dict[int, str], template: TemplatePlanilha) -> dict[str, Any]:
    registro: dict[str, Any] = {}
    for i, campo in posicoes.items():
        valor = _texto(linha[i]) if i < len(linha) else None
        if valor is not None and template.decimal_brasileiro and campo in template.colunas_decimais:
            valor = decimal_br(valor)
        registro[campo] = valor
    return registro


def _texto(celula: Any) -> str | None:
    """Célula como texto limpo; vazia vira None."""
    if celula is None:
        return None
    texto = str(celula).strip()
    return texto or None


def _linhas_csv(caminho: Path, template: TemplatePlanilha) -> Iterator[list[str]]:
    with caminho.open(encoding=template.encoding, newline="") as arquivo:
        leitor = csv.reader(arquivo, delimiter=template.delimitador)
        for numero, linha in enumerate(leitor, start=1):
            if numero >= template.linha_cabecalho:
                yield linha


def _linhas_xlsx(caminho: Path, template: TemplatePlanilha) -> Iterator[list[Any]]:
    from openpyxl import load_workbook  # import tardio: conector de CSV não precisa

    livro = load_workbook(caminho, read_only=True, data_only=True)
    try:
        planilha = livro[template.aba] if template.aba else livro[livro.sheetnames[0]]
        for numero, linha in enumerate(planilha.iter_rows(values_only=True), start=1):
            if numero >= template.linha_cabecalho:
                yield list(linha)
    finally:
        livro.close()
